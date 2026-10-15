"""Tests for the ATS CV Audit & Rebuild ("CV Doctor") pipeline."""
import io
import os
import re

from tests.conftest import register_and_login

FIXTURE = os.path.join(os.path.dirname(__file__), "fixtures", "sample_cv_operations.txt")
SOURCE = open(FIXTURE, encoding="utf8").read()


def _auth(tokens):
    return {"Authorization": f"Bearer {tokens['access_token']}"}


def _messy_docx(text: str = SOURCE) -> tuple:
    """A deliberately ATS-unfriendly Word CV: contact details in the page
    header and the skills laid out in a table."""
    from docx import Document
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "sipho.nkosi@example.com | 082 555 0199"
    for line in text.splitlines():
        if line.startswith("Process monitoring"):
            t = doc.add_table(rows=1, cols=3)
            for i, cell in enumerate(line.split(",")[:3]):
                t.rows[0].cells[i].text = cell
            continue
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return ("Sipho CV.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


# ---- Unit: parser ------------------------------------------------------------------

def test_parser_extracts_roles_dates_and_statuses():
    from app.documents.cv_audit.parse import parse_cv_text
    cv = parse_cv_text(SOURCE)
    assert cv["full_name"] == "Sipho James Nkosi"
    assert cv["email"] == "sipho.nkosi@example.com"
    roles = [(e["position"], e["employer"], e["start"], e["end"]) for e in cv["experience"]]
    assert roles == [
        ("Process Controller Assistant", "Rand Water", "Mar 2022", "Present"),
        ("Rolls Plant Process Controller", "Barloworld Ingrain SA", "Jan 2019", "Feb 2022"),
        ("Microbiology Laboratory Technician", "Tongaat Hulett Starch", "Jun 2016", "Dec 2018"),
    ]
    status = {e["qualification"]: e["status"] for e in cv["education"]}
    assert status["Postgraduate Diploma in Operations Management"] == "awaiting_conferral"
    assert status["National Diploma: Biotechnology"] == "completed"
    certs = {c["name"]: c for c in cv["certifications"]}
    assert certs["Water and Wastewater Process Control NQF 4"]["status"] == "in_progress"
    assert certs["National Certificate: Chemical Operations"]["issuer"] == "CHIETA"
    kinds = {s["kind"] for s in cv["sensitive"]}
    assert {"id_number", "date_of_birth", "gender", "marital_status"} <= kinds
    assert cv["references_present"] is True
    assert cv["phone"] == "082 555 0199"


def test_parser_repairs_pdf_bullet_artefacts():
    from app.documents.cv_audit.parse import parse_cv_text
    text = ("Jane Dube\njane@example.com\nWORK EXPERIENCE\nLab Analyst\nAcme Foods (Pty) Ltd | Jan 2020 – Present\n"
            "\x7f\n\nTested raw materials against the product\nspecification every shift\n Calibrated balances\n")
    exp = parse_cv_text(text)["experience"]
    assert len(exp) == 1
    assert exp[0]["bullets"] == ["Tested raw materials against the product specification every shift",
                                 "Calibrated balances"]


# ---- Unit: rewrite truthfulness ----------------------------------------------------------

def test_rebuild_never_invents_numbers_or_upgrades_verbs():
    from app.documents.cv_audit.parse import parse_cv_text
    from app.documents.cv_audit.rewrite import build_rebuilt_cv
    cv, _ = build_rebuilt_cv(parse_cv_text(SOURCE), SOURCE)
    source_numbers = set(re.findall(r"\d+", SOURCE))
    produced = " ".join([cv["summary"], cv["headline"]] + [b for e in cv["experience"] for b in e["bullets"]])
    assert set(re.findall(r"\d+", produced)) <= source_numbers
    bullets = [b for e in cv["experience"] for b in e["bullets"]]
    # "Assisted with…" / "Helped with…" stay assistance — never become "Led"/"Managed".
    assert "Assist with chemical dosing adjustments under supervision" in bullets
    assert "Assisted with training of 3 new operators" in bullets
    assert not any(b.split()[0] in ("Led", "Managed", "Supervised", "Spearheaded") for b in bullets)
    # Current role in present tense, past roles in past tense.
    assert cv["experience"][0]["bullets"][0] == "Monitor sedimentation tanks and filters"
    assert "Conducted microbiological testing of raw materials and products" in cv["experience"][2]["bullets"]
    # Titles, employers and qualification names are untouched.
    assert [e["position"] for e in cv["experience"]][1] == "Rolls Plant Process Controller"
    assert cv["current_studies"][0]["qualification"] == "Postgraduate Diploma in Operations Management"
    # Sensitive data never reaches the rebuilt CV.
    blob = repr({k: v for k, v in cv.items() if k != "references"})
    assert "9001015009087" not in blob and "Married" not in blob
    # References are the South African standard: kept, and the referee's number is never
    # mistaken for the candidate's own.
    assert cv["references"] == [{"name": "Mr T. Mokoena", "details": "Shift Supervisor",
                                 "phone": "016 555 0100", "email": None}]
    assert cv["phone"] == "082 555 0199"
    # Acronyms spelt out once, in reading order (competencies come before experience).
    comps = [t for ts in cv["competencies"].values() for t in ts]
    assert "Supervisory Control and Data Acquisition (SCADA)" in comps
    assert any("SCADA" in b and "Supervisory" not in b for b in bullets)


def test_competencies_are_all_evidenced():
    from app.documents.cv_audit.parse import parse_cv_text
    from app.documents.cv_audit.rewrite import build_rebuilt_cv
    cv, _ = build_rebuilt_cv(parse_cv_text(SOURCE), SOURCE)
    low = SOURCE.lower()
    for term in [t for ts in cv["competencies"].values() for t in ts]:
        core = term.split(" (")[-1].rstrip(")") if "(" in term else term
        words = [w for w in re.findall(r"[a-z]{4,}", core.lower())]
        assert any(w[:5] in low for w in words) or core.lower() in low, term


# ---- End to end ---------------------------------------------------------------------------

def _upload_and_audit(client, tokens):
    up = client.post("/api/v1/cv", headers=_auth(tokens), files={"file": _messy_docx()})
    assert up.status_code == 201, up.text
    resp = client.post("/api/v1/cv-audits", headers=_auth(tokens), json={"cv_id": up.json()["id"]})
    assert resp.status_code == 201, resp.text
    return up.json(), resp.json()


def test_full_audit_flow(client):
    _, tokens = register_and_login(client)
    cv, audit = _upload_and_audit(client, tokens)

    codes = {f["code"] for f in audit["findings"]}
    assert {"tables", "contact_in_header"} <= codes
    before, after = audit["original_scores"], audit["final_scores"]
    assert after["areas"]["ats_compatibility"]["score"] > before["areas"]["ats_compatibility"]["score"]
    assert after["overall"] > before["overall"]
    assert "not an official rating" in after["disclaimer"]
    assert audit["stress"]["passed"] == audit["stress"]["total"], audit["stress"]["checks"]
    assert audit["pages"] <= 2
    assert any(i["category"] == "Achievements" for i in audit["checklist"])
    assert 1 <= len(audit["directions"]) <= 3

    from pypdf import PdfReader
    from docx import Document
    for kind, check in (("pdf", None), ("docx", None), ("report", None)):
        r = client.get(f"/api/v1/cv-audits/{audit['id']}/file/{kind}", headers=_auth(tokens))
        assert r.status_code == 200, (kind, r.text)
        assert "attachment" in r.headers["content-disposition"]
        if kind == "docx":
            d = Document(io.BytesIO(r.content))
            assert not d.tables
            assert not "".join(p.text for p in d.sections[0].header.paragraphs).strip()
            assert d.paragraphs[0].text == "Sipho James Nkosi"
        else:
            text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(r.content)).pages)
            assert ("Sipho James Nkosi" in text) if kind == "pdf" else ("ATS CV Audit Report" in text)

    # Candidate adds a verified figure and confirms a role; everything re-scores.
    rebuilt = audit["rebuilt"]
    rebuilt["experience"][0]["bullets"].append("Monitor 12 sedimentation tanks per shift")
    r = client.put(f"/api/v1/cv-audits/{audit['id']}/cv", headers=_auth(tokens), json=rebuilt)
    assert r.status_code == 200, r.text
    edited = r.json()
    assert edited["edited_by_candidate"] is True
    assert "Monitor 12 sedimentation tanks per shift" in edited["rebuilt"]["experience"][0]["bullets"]
    assert edited["final_scores"]["areas"]["recruitment_effectiveness"]["score"] >= \
        after["areas"]["recruitment_effectiveness"]["score"]

    lst = client.get("/api/v1/cv-audits", headers=_auth(tokens)).json()
    assert [a["id"] for a in lst] == [audit["id"]]


def test_audit_is_private_and_deleted_with_cv(client):
    _, owner = register_and_login(client)
    cv, audit = _upload_and_audit(client, owner)
    _, other = register_and_login(client, email="other@example.com")
    for path in (f"/api/v1/cv-audits/{audit['id']}", f"/api/v1/cv-audits/{audit['id']}/file/pdf"):
        assert client.get(path, headers=_auth(other)).status_code == 404
    assert client.post("/api/v1/cv-audits", headers=_auth(other), json={"cv_id": cv["id"]}).status_code == 404
    assert client.delete(f"/api/v1/cv/{cv['id']}", headers=_auth(owner)).status_code == 204
    assert client.get(f"/api/v1/cv-audits/{audit['id']}", headers=_auth(owner)).status_code == 404


def test_edit_payload_is_capped(client):
    _, tokens = register_and_login(client)
    _, audit = _upload_and_audit(client, tokens)
    bad = audit["rebuilt"]
    bad["summary"] = "x" * 10000
    r = client.put(f"/api/v1/cv-audits/{audit['id']}/cv", headers=_auth(tokens), json=bad)
    assert r.status_code == 422


def test_unreadable_cv_is_refused(client):
    _, tokens = register_and_login(client)
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.rect(50, 50, 200, 200, fill=1)  # an "image-only" page with no text
    c.save()
    buf.seek(0)
    up = client.post("/api/v1/cv", headers=_auth(tokens), files={"file": ("scan.pdf", buf, "application/pdf")})
    assert up.status_code == 201
    r = client.post("/api/v1/cv-audits", headers=_auth(tokens), json={"cv_id": up.json()["id"]})
    assert r.status_code == 409


def test_apply_rebuilt_cv_to_profile_is_add_only_and_keeps_status(client):
    _, tokens = register_and_login(client)
    h = _auth(tokens)
    # A role the candidate already confirmed by hand must never be touched.
    client.post("/api/v1/profile/experience", headers=h, json={
        "employer": "Barloworld Ingrain SA", "position": "Rolls Plant Process Controller",
        "responsibilities": "My own wording."})
    client.post("/api/v1/profile/skills", headers=h, json={"name": "Chemical dosing", "category": "technical"})
    _, audit = _upload_and_audit(client, tokens)

    r = client.post(f"/api/v1/cv-audits/{audit['id']}/apply-to-profile", headers=h, json={})
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["work_experience"] == 2 and res["skipped_duplicates"] >= 2
    assert res["education"] == 4 and res["certifications"] == 3
    assert "current_occupation" in res["profile_fields"]

    facts = client.get("/api/v1/profile/master-cv", headers=h).json()
    work = {w["employer"]: w for w in facts["experience"]}
    assert work["Barloworld Ingrain SA"]["responsibilities"] == "My own wording."
    rand = work["Rand Water"]
    assert rand["is_current"] is True and rand["start_date"] == "2022-03-01"
    assert rand["confirmed_by_candidate"] is False
    quals = [e["qualification"] for e in facts["education"]]
    assert "Postgraduate Diploma in Operations Management (completed – awaiting conferral)" in quals
    assert "National Diploma: Biotechnology" in quals
    assert [s for s in facts["skills"] if s.lower() == "chemical dosing"] == ["Chemical dosing"]

    # Applying twice adds nothing new.
    again = client.post(f"/api/v1/cv-audits/{audit['id']}/apply-to-profile", headers=h, json={}).json()
    assert again["skills"] == again["work_experience"] == again["education"] == again["certifications"] == 0

    _, other = register_and_login(client, email="nosy@example.com")
    assert client.post(f"/api/v1/cv-audits/{audit['id']}/apply-to-profile", headers=_auth(other),
                       json={}).status_code == 404


# ---- ATS compliance guarantee ------------------------------------------------------------

_HOSTILE_BITS = ["🚀 ", "★★★☆☆ ", "\t", "  ", "• ", " ", "➤ ", "<b>", "&", "​", "\x07", "✔ ",
                 "“quoted”", "R12 500", "x" * 60, "Ñandú – São Tomé", "isiZulu", "50%"]
_ODD_DATES = ["03/2021", "January 2020", "2019/05", "2018", "Sept 2017", "present", "", None, "to date"]


def _hostile_cv(rng) -> dict:
    def txt(base: str) -> str:
        return base + "".join(rng.choice(_HOSTILE_BITS) for _ in range(rng.randint(0, 3)))

    roles = []
    for i in range(rng.randint(0, 14)):
        current = i == 0 and rng.random() < 0.6
        roles.append({
            "position": txt(rng.choice(["Process Controller", "Lab Analyst", "Ops Lead", ""])) or None,
            "employer": txt(rng.choice(["Rand Water", "Acme (Pty) Ltd", "Société Générale", ""])) or None,
            "location": rng.choice([None, "Vereeniging, Gauteng", "Paris"]),
            "start": rng.choice(_ODD_DATES), "end": rng.choice(_ODD_DATES), "is_current": current,
            "bullets": [txt(rng.choice(["Monitored tanks.", "- Ran SCADA", "• Did QA & QC", "x" * 500]))
                        for _ in range(rng.randint(0, 25))],
        })
    edu = [{"qualification": txt("National Diploma: Biotechnology"), "institution": rng.choice(["VUT", None]),
            "date_text": rng.choice(["2016", "Jan 2014 - Dec 2016", "03/2021", None]),
            "status": rng.choice(["completed", "in_progress", "awaiting_conferral", "unverified"])}
           for _ in range(rng.randint(0, 5))]
    return {
        "full_name": rng.choice([txt("Nomvula Mthembu-Ndlovu"), None, "José da Silva 😀"]),
        "headline": txt("Process Controller | Water"),
        "email": rng.choice(["n.mthembu@example.com", None]), "phone": rng.choice(["082 555 0199", None]),
        "location": rng.choice(["Vereeniging, Gauteng", None]),
        "summary": txt("Experienced operator. " * rng.randint(0, 40)),
        "competencies": {txt("Water & Wastewater"): [txt("Jar testing"), txt("SCADA"), txt("Chemical dosing")]
                         for _ in range(rng.randint(0, 2))},
        "software": [txt("Microsoft Excel")], "experience": roles, "education": edu,
        "current_studies": edu[:1], "certifications": [{"name": txt("First Aid Level 1"), "issuer": None,
                                                         "year": "2023", "kind": "short_course",
                                                         "status": "completed"}],
        "languages": [txt("English"), "isiZulu"], "drivers_licence": rng.choice([None, "Code 10 (C1)"]),
        "projects": [txt("Commissioned filter 5")], "achievements": [],
        "references": [{"name": txt("Mr T. Mokoena"), "details": txt("Shift Supervisor, Rand Water"),
                        "phone": rng.choice(["016 555 0100", None]), "email": rng.choice([None, "t@rw.co.za"])}
                       for _ in range(rng.randint(0, 4))],
    }


def test_every_rebuilt_cv_is_ats_compliant_even_with_hostile_input():
    """Whatever the parser found or the candidate typed — emoji, rating stars, Word
    symbol bullets, tabs, HTML-ish text, odd date formats, 14 roles with 25 bullets,
    missing fields — the generated files must pass every ATS-format check."""
    import random
    from app.documents.cv_audit.stress import stress_test
    rng = random.Random(20261015)
    for case in range(30):
        cv = _hostile_cv(rng)
        result = stress_test(cv)
        failed = [c for c in result["checks"] if not c["ok"]]
        assert result["ats_compliant"], (case, failed)
        assert not re.search(r"[\U0001f000-\U0001faff★☆\x07​]", result["pdf_text"]), case


def test_missing_contact_is_a_content_blocker_not_invented():
    from app.documents.cv_audit.parse import parse_cv_text
    from app.documents.cv_audit.rewrite import build_rebuilt_cv
    from app.documents.cv_audit.stress import stress_test
    text = SOURCE.replace("Cell: 082 555 0199 | Email: sipho.nkosi@example.com", "")
    cv, _ = build_rebuilt_cv(parse_cv_text(text), text)
    result = stress_test(cv)
    assert result["ats_compliant"] is True
    assert result["ats_ready"] is False
    fields = {b["field"] for b in result["content_blockers"]}
    assert {"email", "phone"} <= fields
    assert "@" not in result["pdf_text"]


def test_references_parse_in_common_south_african_layouts():
    from app.documents.cv_audit.parse import _parse_references
    refs, _ = _parse_references([
        "1. John Smith – Production Manager, ABC Foods – 082 123 4567 – john@abc.co.za",
        "2. Lerato Mokoena, HR Officer, Rand Water, 016 421 0000"])
    assert [(r["name"], r["phone"]) for r in refs] == [("John Smith", "082 123 4567"),
                                                        ("Lerato Mokoena", "016 421 0000")]
    refs, _ = _parse_references(["Name: Mrs P Dlamini", "Position: Plant Manager", "Company: Sasol",
                                 "Tel: 011 555 1234", "Name: Mr J Naidoo", "Email: j.naidoo@example.com"])
    assert [r["name"] for r in refs] == ["Mrs P Dlamini", "Mr J Naidoo"]
    assert refs[0]["details"] == "Plant Manager, Sasol"
    assert _parse_references(["Available on request"]) == ([], True)
    # Stray text under the heading (a DOCX skills table extracted at the end) is not a referee.
    refs, _ = _parse_references(["Mr T. Mokoena, Shift Supervisor, 016 555 0100",
                                 "Process monitoring ★★★★☆  SCADA  Jar testing"])
    assert [r["name"] for r in refs] == ["Mr T. Mokoena"]


def test_references_render_and_missing_ones_are_flagged():
    from app.documents.cv_audit.parse import parse_cv_text
    from app.documents.cv_audit.rewrite import build_rebuilt_cv
    from app.documents.cv_audit.directions import build_checklist
    from app.documents.cv_audit.stress import stress_test
    parsed = parse_cv_text(SOURCE)
    cv, _ = build_rebuilt_cv(parsed, SOURCE)
    result = stress_test(cv)
    assert result["ats_compliant"]
    tail = result["pdf_text"][result["pdf_text"].find("REFERENCES"):]
    assert "Mr T. Mokoena" in tail and "016 555 0100" in tail
    items = [i["item"] for i in build_checklist(parsed, cv) if i["category"] == "References"]
    assert any("one or two more" in i for i in items)   # only one referee given
